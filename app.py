import os
import sqlite3
from datetime import date, datetime
from functools import wraps
from io import BytesIO
from typing import Dict, List, Tuple

from flask import (
    Flask,
    Response,
    flash,
    g,
    jsonify,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook, load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from werkzeug.utils import secure_filename

EXPENSE_TYPES: List[str] = ["Materialy", "Zaliczka", "Usluga", "Robocizna", "Inne"]
CATEGORIES: List[str] = [f"ETAP {i}" for i in range(6)]
ROLE_LABELS: Dict[str, str] = {"admin": "Admin", "user": "Użytkownik", "guest": "Gość"}
ALLOWED_ATTACHMENT_EXTENSIONS: Tuple[str, ...] = (
    ".pdf",
    ".png",
    ".jpg",
    ".jpeg",
    ".heic",
    ".heif",
    ".webp",
)
MAX_ATTACHMENT_SIZE = 10 * 1024 * 1024  # 10 MB


def create_app() -> Flask:
    app = Flask(__name__)
    app.config.update(
        SECRET_KEY=os.environ.get("SECRET_KEY", "super-secret-key"),
        DATABASE=os.path.join(app.root_path, "expenses.db"),
        USERNAME=os.environ.get("APP_USERNAME", "lukasz"),
        PASSWORD=os.environ.get("APP_PASSWORD", "lukasz29"),
        MAX_ATTACHMENT_SIZE=int(os.environ.get("MAX_ATTACHMENT_SIZE", MAX_ATTACHMENT_SIZE)),
    )

    os.makedirs(app.root_path, exist_ok=True)

    def get_db() -> sqlite3.Connection:
        if "db" not in g:
            g.db = sqlite3.connect(app.config["DATABASE"])
            g.db.row_factory = sqlite3.Row
        return g.db

    def close_db(_: Exception | None) -> None:
        db = g.pop("db", None)
        if db is not None:
            db.close()

    def ensure_schema() -> None:
        db = get_db()
        db.execute(
            """
            CREATE TABLE IF NOT EXISTS expenses (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                expense_date TEXT NOT NULL,
                merchant TEXT NOT NULL,
                amount REAL NOT NULL,
                category TEXT NOT NULL,
                notes TEXT,
                bank TEXT,
                description TEXT,
                expense_type TEXT NOT NULL,
                display_order INTEGER NOT NULL DEFAULT 0
            )
            """
        )
        expense_columns = {
            row["name"] for row in db.execute("PRAGMA table_info(expenses)").fetchall()
        }
        if "display_order" not in expense_columns:
            db.execute(
                "ALTER TABLE expenses ADD COLUMN display_order INTEGER NOT NULL DEFAULT 0"
            )
            db.execute("UPDATE expenses SET display_order = id WHERE display_order = 0")

        if "attachment_filename" not in expense_columns:
            db.execute("ALTER TABLE expenses ADD COLUMN attachment_filename TEXT")
        if "attachment_mimetype" not in expense_columns:
            db.execute("ALTER TABLE expenses ADD COLUMN attachment_mimetype TEXT")
        if "attachment_data" not in expense_columns:
            db.execute("ALTER TABLE expenses ADD COLUMN attachment_data BLOB")

        db.execute(
            """
            CREATE TABLE IF NOT EXISTS users (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                username TEXT UNIQUE NOT NULL,
                password TEXT NOT NULL,
                role TEXT NOT NULL,
                can_add INTEGER NOT NULL DEFAULT 1,
                can_edit INTEGER NOT NULL DEFAULT 1,
                can_delete INTEGER NOT NULL DEFAULT 1,
                can_view_reports INTEGER NOT NULL DEFAULT 1
            )
            """
        )

        ensure_default_admin(db)
        db.commit()

    def ensure_default_admin(db: sqlite3.Connection) -> None:
        existing = db.execute(
            "SELECT id FROM users WHERE username = ?", (app.config["USERNAME"],)
        ).fetchone()
        if existing is None:
            db.execute(
                """
                INSERT INTO users (
                    username, password, role, can_add, can_edit, can_delete, can_view_reports
                ) VALUES (?, ?, 'admin', 1, 1, 1, 1)
                """,
                (app.config["USERNAME"], app.config["PASSWORD"]),
            )

    @app.before_request
    def before_request() -> None:
        ensure_schema()
        g.user = None
        g.permissions = {}
        user_id = session.get("user_id")
        if user_id is not None:
            user = (
                get_db()
                .execute("SELECT * FROM users WHERE id = ?", (user_id,))
                .fetchone()
            )
            if user is not None:
                g.user = user
                g.permissions = {
                    "can_add": bool(user["can_add"]),
                    "can_edit": bool(user["can_edit"]),
                    "can_delete": bool(user["can_delete"]),
                    "can_view_reports": bool(user["can_view_reports"]),
                }
            else:
                session.clear()

    @app.teardown_appcontext
    def teardown_db(exception: Exception | None) -> None:
        close_db(exception)

    def login_required(view):
        @wraps(view)
        def wrapped_view(**kwargs):
            if g.get("user") is None:
                return redirect(url_for("login"))
            return view(**kwargs)

        return wrapped_view

    def permission_required(permission: str):
        def decorator(view):
            @wraps(view)
            def wrapped_view(**kwargs):
                if g.get("user") is None:
                    return redirect(url_for("login"))
                if not g.permissions.get(permission, False):
                    if request.accept_mimetypes.accept_json:
                        return jsonify({"error": "Brak uprawnień."}), 403
                    flash("Brak uprawnień do wykonania tej operacji.", "error")
                    return redirect(url_for("index"))
                return view(**kwargs)

            return wrapped_view

        return decorator

    def admin_required(view):
        @wraps(view)
        def wrapped_view(**kwargs):
            if g.get("user") is None:
                return redirect(url_for("login"))
            if g.user["role"] != "admin":
                flash("Ta sekcja wymaga uprawnień administratora.", "error")
                return redirect(url_for("index"))
            return view(**kwargs)

        return wrapped_view

    def parse_amount(raw_value: str) -> float:
        cleaned = (raw_value or "").replace(" ", "").replace(",", ".")
        try:
            return float(cleaned)
        except ValueError as exc:
            raise ValueError(
                "Nieprawidłowa kwota. Użyj cyfr oraz kropki lub przecinka."
            ) from exc

    def prepare_attachment(file_storage) -> Tuple[str, str, bytes] | None:
        if not file_storage or not file_storage.filename:
            return None

        filename = secure_filename(file_storage.filename)
        if not filename:
            raise ValueError("Niepoprawna nazwa pliku. Spróbuj ponownie.")

        extension = os.path.splitext(filename)[1].lower()
        if extension not in ALLOWED_ATTACHMENT_EXTENSIONS:
            allowed = ", ".join(ext.lstrip(".") for ext in ALLOWED_ATTACHMENT_EXTENSIONS)
            raise ValueError(f"Nieobsługiwany format pliku. Dozwolone rozszerzenia: {allowed}.")

        file_storage.stream.seek(0)
        data = file_storage.read()
        max_size = app.config.get("MAX_ATTACHMENT_SIZE", MAX_ATTACHMENT_SIZE)
        if len(data) > max_size:
            raise ValueError(
                "Załącznik jest zbyt duży. Maksymalny rozmiar pliku to 10 MB."
            )

        mimetype = file_storage.mimetype or "application/octet-stream"
        return filename, mimetype, data

    def fetch_category_totals(db: sqlite3.Connection) -> Dict[str, float]:
        results = db.execute(
            "SELECT category, SUM(amount) as total FROM expenses GROUP BY category"
        ).fetchall()
        return {row["category"]: row["total"] for row in results}

    def fetch_type_totals(db: sqlite3.Connection) -> Dict[str, float]:
        results = db.execute(
            "SELECT expense_type, SUM(amount) as total FROM expenses GROUP BY expense_type"
        ).fetchall()
        return {row["expense_type"]: row["total"] for row in results}

    def fetch_monthly_totals(db: sqlite3.Connection) -> List[sqlite3.Row]:
        return db.execute(
            """
            SELECT substr(expense_date, 1, 7) AS month, SUM(amount) AS total
            FROM expenses
            GROUP BY substr(expense_date, 1, 7)
            ORDER BY month
            """
        ).fetchall()

    def fetch_expenses(
        db: sqlite3.Connection, requested_sort: str, direction: str
    ) -> tuple[list[sqlite3.Row], str, str]:
        sort_map = {
            "display_order": "display_order",
            "date": "expense_date",
            "amount": "amount",
            "category": "category",
            "type": "expense_type",
            "merchant": "merchant",
        }
        sort_column = sort_map.get(requested_sort, "display_order")
        normalized_direction = direction.lower()
        if sort_column == "display_order":
            normalized_direction = "asc"
        elif normalized_direction not in {"asc", "desc"}:
            normalized_direction = "asc"

        rows = db.execute(
            f"SELECT * FROM expenses ORDER BY {sort_column} {normalized_direction.upper()}, id ASC"
        ).fetchall()
        return rows, sort_column, normalized_direction

    def parse_excel_date(value) -> str:
        if value is None:
            raise ValueError("Brak daty")
        if isinstance(value, datetime):
            return value.date().isoformat()
        if isinstance(value, date):
            return value.isoformat()
        if isinstance(value, str):
            stripped = value.strip()
            for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y", "%Y/%m/%d"):
                try:
                    return datetime.strptime(stripped, fmt).date().isoformat()
                except ValueError:
                    continue
            raise ValueError("Nieprawidłowy format daty")
        raise ValueError("Nieobsługiwany format daty")

    @app.context_processor
    def inject_globals() -> Dict[str, object]:
        return {
            "current_user": g.get("user"),
            "permissions": g.get("permissions", {}),
            "role_labels": ROLE_LABELS,
            "expense_types": EXPENSE_TYPES,
            "categories": CATEGORIES,
        }

    @app.route("/login", methods=["GET", "POST"])
    def login():
        if request.method == "POST":
            username = request.form.get("username", "").strip()
            password = request.form.get("password", "")
            user = (
                get_db()
                .execute("SELECT * FROM users WHERE username = ?", (username,))
                .fetchone()
            )
            if user and user["password"] == password:
                session.clear()
                session["logged_in"] = True
                session["user_id"] = user["id"]
                flash("Zalogowano pomyślnie.", "success")
                return redirect(url_for("index"))
            flash("Niepoprawny login lub hasło", "error")
        return render_template("login.html")

    @app.route("/logout")
    def logout():
        session.clear()
        flash("Zostałeś poprawnie wylogowany.", "success")
        return redirect(url_for("login"))

    @app.route("/")
    @login_required
    def index():
        db = get_db()
        expenses_rows, _, _ = fetch_expenses(db, "display_order", "asc")

        category_totals = fetch_category_totals(db)
        type_totals = fetch_type_totals(db)
        total_cost = sum(row["amount"] for row in expenses_rows)

        return render_template(
            "index.html",
            category_totals=category_totals,
            type_totals=type_totals,
            total_cost=total_cost,
        )

    @app.route("/history")
    @login_required
    def history():
        db = get_db()
        requested_sort = request.args.get("sort", "display_order")
        direction = request.args.get("direction", "asc")
        expenses_rows, _, normalized_direction = fetch_expenses(
            db, requested_sort, direction
        )
        expenses: List[Dict[str, object]] = []
        for row in expenses_rows:
            expense = dict(row)
            has_attachment = bool(expense.get("attachment_data"))
            expense["has_attachment"] = has_attachment
            expense["attachment_url"] = (
                url_for("download_attachment", expense_id=expense["id"])
                if has_attachment
                else ""
            )
            expense.pop("attachment_data", None)
            expenses.append(expense)
        return render_template(
            "history.html",
            expenses=expenses,
            current_sort=requested_sort,
            current_direction=normalized_direction,
        )

    @app.post("/expenses")
    @login_required
    @permission_required("can_add")
    def add_expense():
        expense_date = request.form.get("expense_date", "").strip()
        merchant = request.form.get("merchant", "").strip()
        amount_raw = request.form.get("amount", "")
        category = request.form.get("category", "").strip()
        notes = request.form.get("notes", "").strip()
        bank = request.form.get("bank", "").strip()
        description = request.form.get("description", "").strip()
        expense_type = request.form.get("expense_type", "").strip()
        attachment_file = request.files.get("attachment")

        try:
            amount = parse_amount(amount_raw)
        except ValueError as exc:
            flash(str(exc), "error")
            return redirect(url_for("index"))

        try:
            attachment_payload = prepare_attachment(attachment_file)
        except ValueError as exc:
            flash(str(exc), "error")
            return redirect(url_for("index"))

        attachment_filename = attachment_mimetype = None
        attachment_data = None
        if attachment_payload:
            attachment_filename, attachment_mimetype, attachment_data = attachment_payload

        if (
            not expense_date
            or not merchant
            or not category
            or not expense_type
            or expense_type not in EXPENSE_TYPES
            or category not in CATEGORIES
        ):
            flash("Proszę uzupełnić wymagane pola poprawnymi wartościami.", "error")
            return redirect(url_for("index"))

        db = get_db()
        max_order_row = db.execute(
            "SELECT COALESCE(MAX(display_order), -1) AS max_order FROM expenses"
        ).fetchone()
        current_max = max_order_row["max_order"] if max_order_row is not None else -1
        next_order = current_max + 1
        db.execute(
            """
            INSERT INTO expenses (
                expense_date, merchant, amount, category, notes, bank, description,
                expense_type, display_order, attachment_filename, attachment_mimetype, attachment_data
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                expense_date,
                merchant,
                amount,
                category,
                notes,
                bank,
                description,
                expense_type,
                next_order,
                attachment_filename,
                attachment_mimetype,
                attachment_data,
            ),
        )
        db.commit()
        flash("Wydatek został dodany.", "success")
        return redirect(url_for("index"))

    @app.post("/expenses/<int:expense_id>/update")
    @login_required
    @permission_required("can_edit")
    def update_expense(expense_id: int):
        expense_date = request.form.get("expense_date", "").strip()
        merchant = request.form.get("merchant", "").strip()
        amount_raw = request.form.get("amount", "")
        category = request.form.get("category", "").strip()
        notes = request.form.get("notes", "").strip()
        bank = request.form.get("bank", "").strip()
        description = request.form.get("description", "").strip()
        expense_type = request.form.get("expense_type", "").strip()
        is_ajax = request.headers.get("X-Requested-With") == "XMLHttpRequest"
        remove_attachment = request.form.get("remove_attachment") in {"on", "1", "true"}
        attachment_file = request.files.get("attachment")

        try:
            amount = parse_amount(amount_raw)
        except ValueError as exc:
            message = str(exc)
            flash(message, "error")
            if is_ajax:
                return jsonify({"error": message}), 400
            return redirect(url_for("history"))

        if (
            not expense_date
            or not merchant
            or not category
            or not expense_type
            or expense_type not in EXPENSE_TYPES
            or category not in CATEGORIES
        ):
            message = "Proszę uzupełnić wymagane pola poprawnymi wartościami."
            flash(message, "error")
            if is_ajax:
                return jsonify({"error": message}), 400
            return redirect(url_for("history"))

        db = get_db()
        existing = db.execute(
            "SELECT attachment_filename, attachment_mimetype, attachment_data FROM expenses WHERE id = ?",
            (expense_id,),
        ).fetchone()
        if existing is None:
            message = "Nie znaleziono wpisu do edycji."
            if is_ajax:
                return jsonify({"error": message}), 404
            flash(message, "error")
            return redirect(url_for("history"))

        current_attachment = {
            "filename": existing["attachment_filename"],
            "mimetype": existing["attachment_mimetype"],
            "data": existing["attachment_data"],
        }

        try:
            attachment_payload = prepare_attachment(attachment_file)
        except ValueError as exc:
            message = str(exc)
            if is_ajax:
                return jsonify({"error": message}), 400
            flash(message, "error")
            return redirect(url_for("history"))

        if attachment_payload:
            attachment_filename, attachment_mimetype, attachment_data = attachment_payload
        else:
            attachment_filename = current_attachment["filename"]
            attachment_mimetype = current_attachment["mimetype"]
            attachment_data = current_attachment["data"]

        if remove_attachment:
            attachment_filename = None
            attachment_mimetype = None
            attachment_data = None

        db.execute(
            """
            UPDATE expenses
               SET expense_date = ?,
                   merchant = ?,
                   amount = ?,
                   category = ?,
                   notes = ?,
                   bank = ?,
                   description = ?,
                   expense_type = ?,
                   attachment_filename = ?,
                   attachment_mimetype = ?,
                   attachment_data = ?
             WHERE id = ?
            """,
            (
                expense_date,
                merchant,
                amount,
                category,
                notes,
                bank,
                description,
                expense_type,
                attachment_filename,
                attachment_mimetype,
                attachment_data,
                expense_id,
            ),
        )
        db.commit()
        success_message = "Wydatek został zaktualizowany."
        flash(success_message, "success")
        if is_ajax:
            return jsonify({"redirect": url_for("history")})
        return redirect(url_for("history"))

    @app.post("/expenses/<int:expense_id>/delete")
    @login_required
    @permission_required("can_delete")
    def delete_expense(expense_id: int):
        db = get_db()
        db.execute("DELETE FROM expenses WHERE id = ?", (expense_id,))
        db.commit()
        flash("Wpis został usunięty.", "success")
        return redirect(url_for("history"))

    @app.post("/expenses/reorder")
    @login_required
    @permission_required("can_edit")
    def reorder_expenses() -> Response:
        payload = request.get_json(silent=True) or {}
        order = payload.get("order", [])
        if not isinstance(order, list):
            return jsonify({"error": "Nieprawidłowe dane."}), 400
        db = get_db()
        try:
            normalized_order = [int(expense_id) for expense_id in order]
        except (TypeError, ValueError):
            return jsonify({"error": "Nieprawidłowe dane."}), 400
        for position, expense_id in enumerate(normalized_order):
            db.execute(
                "UPDATE expenses SET display_order = ? WHERE id = ?",
                (position, expense_id),
            )
        db.commit()
        return ("", 204)

    @app.get("/expenses/<int:expense_id>/attachment")
    @login_required
    def download_attachment(expense_id: int):
        db = get_db()
        row = db.execute(
            """
            SELECT attachment_filename, attachment_mimetype, attachment_data
              FROM expenses
             WHERE id = ?
            """,
            (expense_id,),
        ).fetchone()

        if row is None or row["attachment_data"] is None:
            flash("Brak załącznika do pobrania.", "error")
            return redirect(url_for("history"))

        stream = BytesIO(row["attachment_data"])
        stream.seek(0)
        filename = row["attachment_filename"] or f"zalacznik_{expense_id}"
        mimetype = row["attachment_mimetype"] or "application/octet-stream"
        return send_file(
            stream,
            download_name=filename,
            mimetype=mimetype,
            as_attachment=True,
        )

    @app.get("/expenses/export")
    @login_required
    def export_expenses():
        db = get_db()
        expenses_rows, _, _ = fetch_expenses(db, "display_order", "asc")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Wydatki"
        headers = [
            "Data",
            "Odbiorca",
            "Kwota",
            "Etap",
            "Typ wydatku",
            "Notatki",
            "Bank",
            "Opis",
            "Załącznik",
        ]
        sheet.append(headers)
        for row in expenses_rows:
            sheet.append(
                [
                    row["expense_date"],
                    row["merchant"],
                    float(row["amount"]),
                    row["category"],
                    row["expense_type"],
                    row["notes"] or "",
                    row["bank"] or "",
                    row["description"] or "",
                    "TAK" if row["attachment_filename"] else "",
                ]
            )

        stream = BytesIO()
        workbook.save(stream)
        stream.seek(0)
        filename = "wydatki_budowa.xlsx"
        return send_file(
            stream,
            as_attachment=True,
            download_name=filename,
            mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )

    @app.post("/expenses/import")
    @login_required
    @permission_required("can_add")
    def import_expenses():
        uploaded = request.files.get("file")
        if uploaded is None or uploaded.filename == "":
            flash("Wybierz plik w formacie XLSX.", "error")
            return redirect(url_for("history"))

        try:
            workbook = load_workbook(uploaded, data_only=True)
        except InvalidFileException:
            flash("Nie udało się odczytać pliku. Upewnij się, że to plik .xlsx.", "error")
            return redirect(url_for("history"))

        sheet = workbook.active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            flash("Plik jest pusty.", "error")
            return redirect(url_for("history"))

        header = [str(cell).strip().lower() if cell is not None else "" for cell in rows[0]]
        expected = {
            "data": "expense_date",
            "odbiorca": "merchant",
            "kwota": "amount",
            "etap": "category",
            "typ wydatku": "expense_type",
            "notatki": "notes",
            "bank": "bank",
            "opis": "description",
        }
        header_map = {}
        for label, field in expected.items():
            try:
                index = header.index(label)
            except ValueError:
                flash(
                    "Brakuje kolumny '" + label + "' w nagłówku arkusza.",
                    "error",
                )
                return redirect(url_for("history"))
            header_map[field] = index

        db = get_db()
        max_order_row = db.execute(
            "SELECT COALESCE(MAX(display_order), -1) AS max_order FROM expenses"
        ).fetchone()
        current_max = max_order_row["max_order"] if max_order_row is not None else -1

        inserted = 0
        skipped = 0

        for data_row in rows[1:]:
            if data_row is None:
                continue
            values = {field: data_row[idx] if idx < len(data_row) else None for field, idx in header_map.items()}
            if all(value in (None, "") for value in values.values()):
                continue
            try:
                expense_date = parse_excel_date(values["expense_date"])
                merchant = (values["merchant"] or "").strip()
                expense_type_raw = (values["expense_type"] or "").strip()
                category_raw = (values["category"] or "").strip()
                category = category_raw.upper()
                expense_type = next(
                    (
                        option
                        for option in EXPENSE_TYPES
                        if option.lower() == expense_type_raw.lower()
                    ),
                    expense_type_raw,
                )
                notes = (values["notes"] or "").strip()
                bank = (values["bank"] or "").strip()
                description = (values["description"] or "").strip()
                amount_raw = values["amount"]
                if amount_raw is None or str(amount_raw).strip() == "":
                    raise ValueError("Brak kwoty")
                amount = parse_amount(str(amount_raw))

                if not merchant or category not in CATEGORIES or expense_type not in EXPENSE_TYPES:
                    raise ValueError("Niepoprawne dane w wierszu")

                current_max += 1
                db.execute(
                    """
                    INSERT INTO expenses (
                        expense_date, merchant, amount, category, notes, bank, description,
                        expense_type, display_order, attachment_filename, attachment_mimetype, attachment_data
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """,
                    (
                        expense_date,
                        merchant,
                        amount,
                        category,
                        notes,
                        bank,
                        description,
                        expense_type,
                        current_max,
                        None,
                        None,
                        None,
                    ),
                )
                inserted += 1
            except Exception:
                skipped += 1

        db.commit()
        if inserted:
            flash(
                f"Zaimportowano {inserted} pozycji. Pominięto {skipped} wierszy.",
                "success" if skipped == 0 else "info",
            )
        else:
            flash("Nie udało się zaimportować żadnych danych z pliku.", "error")
        return redirect(url_for("history"))

    @app.route("/users")
    @login_required
    @admin_required
    def manage_users():
        users = [dict(row) for row in get_db().execute("SELECT * FROM users ORDER BY username").fetchall()]
        return render_template("users.html", users=users)

    def count_admins(db: sqlite3.Connection) -> int:
        row = db.execute("SELECT COUNT(*) AS total FROM users WHERE role = 'admin'").fetchone()
        return int(row["total"])

    @app.post("/users/add")
    @login_required
    @admin_required
    def add_user():
        username = request.form.get("username", "").strip()
        password = request.form.get("password", "")
        role = request.form.get("role", "user")
        can_add = 1 if request.form.get("can_add") else 0
        can_edit = 1 if request.form.get("can_edit") else 0
        can_delete = 1 if request.form.get("can_delete") else 0
        can_view_reports = 1 if request.form.get("can_view_reports") else 0

        if not username or not password:
            flash("Nazwa użytkownika i hasło są wymagane.", "error")
            return redirect(url_for("manage_users"))

        if role not in ROLE_LABELS:
            role = "user"

        if role == "admin":
            can_add = can_edit = can_delete = can_view_reports = 1

        db = get_db()
        try:
            db.execute(
                """
                INSERT INTO users (
                    username, password, role, can_add, can_edit, can_delete, can_view_reports
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (username, password, role, can_add, can_edit, can_delete, can_view_reports),
            )
            db.commit()
            flash("Użytkownik został dodany.", "success")
        except sqlite3.IntegrityError:
            flash("Taki użytkownik już istnieje.", "error")
        return redirect(url_for("manage_users"))

    @app.post("/users/<int:user_id>/update")
    @login_required
    @admin_required
    def update_user(user_id: int):
        password = request.form.get("password", "")
        role = request.form.get("role", "user")
        can_add = 1 if request.form.get("can_add") else 0
        can_edit = 1 if request.form.get("can_edit") else 0
        can_delete = 1 if request.form.get("can_delete") else 0
        can_view_reports = 1 if request.form.get("can_view_reports") else 0

        if role not in ROLE_LABELS:
            role = "user"

        db = get_db()
        user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if user is None:
            flash("Nie znaleziono użytkownika.", "error")
            return redirect(url_for("manage_users"))

        if role == "admin":
            can_add = can_edit = can_delete = can_view_reports = 1

        if user["role"] == "admin" and role != "admin" and count_admins(db) <= 1:
            flash("Nie można pozbawić uprawnień ostatniego administratora.", "error")
            return redirect(url_for("manage_users"))

        if not password:
            password = user["password"]

        db.execute(
            """
            UPDATE users
               SET password = ?,
                   role = ?,
                   can_add = ?,
                   can_edit = ?,
                   can_delete = ?,
                   can_view_reports = ?
             WHERE id = ?
            """,
            (password, role, can_add, can_edit, can_delete, can_view_reports, user_id),
        )
        db.commit()
        flash("Dane użytkownika zostały zaktualizowane.", "success")
        return redirect(url_for("manage_users"))

    @app.post("/users/<int:user_id>/delete")
    @login_required
    @admin_required
    def delete_user(user_id: int):
        db = get_db()
        user = db.execute("SELECT * FROM users WHERE id = ?", (user_id,)).fetchone()
        if user is None:
            flash("Nie znaleziono użytkownika.", "error")
            return redirect(url_for("manage_users"))

        if user["role"] == "admin" and count_admins(db) <= 1:
            flash("Nie można usunąć ostatniego administratora.", "error")
            return redirect(url_for("manage_users"))

        if session.get("user_id") == user_id:
            flash("Nie możesz usunąć aktualnie zalogowanego użytkownika.", "error")
            return redirect(url_for("manage_users"))

        db.execute("DELETE FROM users WHERE id = ?", (user_id,))
        db.commit()
        flash("Użytkownik został usunięty.", "success")
        return redirect(url_for("manage_users"))

    @app.route("/reports")
    @login_required
    @permission_required("can_view_reports")
    def reports():
        db = get_db()
        category_totals = fetch_category_totals(db)
        type_totals = fetch_type_totals(db)
        monthly_totals = fetch_monthly_totals(db)

        stage_labels = CATEGORIES
        stage_values = [round(category_totals.get(label, 0.0), 2) for label in stage_labels]

        type_labels = EXPENSE_TYPES
        type_values = [round(type_totals.get(label, 0.0), 2) for label in type_labels]

        month_labels = [row["month"] for row in monthly_totals]
        month_values = [round(row["total"], 2) for row in monthly_totals]

        total_spent = round(sum(stage_values), 2)
        stage_pairs = list(zip(stage_labels, stage_values))
        type_pairs = list(zip(type_labels, type_values))

        top_stage_label, top_stage_value = ("-", 0.0)
        if stage_pairs:
            top_stage_label, top_stage_value = max(stage_pairs, key=lambda item: item[1])

        top_type_label, top_type_value = ("-", 0.0)
        if type_pairs:
            top_type_label, top_type_value = max(type_pairs, key=lambda item: item[1])

        return render_template(
            "reports.html",
            stage_labels=stage_labels,
            stage_values=stage_values,
            type_labels=type_labels,
            type_values=type_values,
            month_labels=month_labels,
            month_values=month_values,
            total_spent=total_spent,
            top_stage_label=top_stage_label,
            top_stage_value=round(top_stage_value, 2),
            top_type_label=top_type_label,
            top_type_value=round(top_type_value, 2),
        )

    return app


app = create_app()


if __name__ == "__main__":
    app.run(debug=True)
